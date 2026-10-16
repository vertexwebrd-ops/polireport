"""Informes descargables (Excel y PDF) y resumen de estadísticas."""
import io
from datetime import datetime, timedelta, timezone
from xml.sax.saxutils import escape

FMT = "%Y-%m-%d %H:%M:%S"


def a_local(s):
    """La base guarda UTC; República Dominicana es UTC-4 todo el año."""
    return (datetime.strptime(s, FMT) - timedelta(hours=4)).strftime("%d/%m/%Y %H:%M") if s else ""


def horas_resolucion(r):
    if not r.get("fecha_resuelto"):
        return None
    return (datetime.strptime(r["fecha_resuelto"], FMT) - datetime.strptime(r["fecha"], FMT)).total_seconds() / 3600


def fmt_horas(h):
    if h is None:
        return "-"
    return f"{h:.1f} h" if h < 24 else f"{h / 24:.1f} días"


def resumen(rows):
    def cuenta(clave):
        d = {}
        for r in rows:
            d[r[clave]] = d.get(r[clave], 0) + 1
        return sorted(d.items(), key=lambda x: -x[1])
    hs = [h for h in (horas_resolucion(r) for r in rows) if h is not None]
    return {"total": len(rows), "por_estado": cuenta("estado"), "por_prioridad": cuenta("prioridad"),
            "por_area": cuenta("lugar"), "por_tipo": cuenta("categoria"),
            "promedio": fmt_horas(sum(hs) / len(hs)) if hs else "-"}


def periodo(desde, hasta):
    if desde and hasta: return f"Del {desde} al {hasta}"
    if desde: return f"Desde el {desde}"
    if hasta: return f"Hasta el {hasta}"
    return "Todos los reportes"


COLUMNAS = [("Fecha", 17, lambda r: a_local(r["fecha"])), ("Área", 30, lambda r: r["lugar"]),
            ("Familia", 22, lambda r: r.get("grupo") or ""), ("Tipo", 18, lambda r: r["categoria"]),
            ("Problema", 32, lambda r: r.get("problema") or ""), ("Urgencia", 10, lambda r: r["prioridad"]),
            ("Estado", 13, lambda r: r["estado"]), ("Descripción", 50, lambda r: r["descripcion"]),
            ("Curso", 18, lambda r: r.get("curso") or ""), ("Reportado por", 22, lambda r: r["autor"]),
            ("Encargado", 20, lambda r: r.get("encargado") or "Sin asignar"),
            ("Personas más", 12, lambda r: r.get("apoyos") or 0),
            ("Resuelto", 17, lambda r: a_local(r["fecha_resuelto"])),
            ("Tiempo de solución", 18, lambda r: fmt_horas(horas_resolucion(r)))]


def xlsx_bytes(rows, desde=None, hasta=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Reportes"
    head = PatternFill("solid", fgColor="173B36")
    ws.append([c[0] for c in COLUMNAS])
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([c[2](r) for c in COLUMNAS])
    for i, c in enumerate(COLUMNAS, 1):
        ws.column_dimensions[get_column_letter(i)].width = c[1]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    rs = wb.create_sheet("Resumen")
    res = resumen(rows)
    rs.append(["PoliReport - Resumen"]); rs["A1"].font = Font(bold=True, size=14)
    rs.append([periodo(desde, hasta)])
    rs.append(["Total de reportes", res["total"]])
    rs.append(["Tiempo promedio de solución", res["promedio"]])
    for titulo, clave in (("Por estado", "por_estado"), ("Por urgencia", "por_prioridad"),
                          ("Por área", "por_area"), ("Por tipo de problema", "por_tipo")):
        rs.append([]); rs.append([titulo, "Reportes"])
        for c in rs[rs.max_row]: c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for nombre, n in res[clave]:
            rs.append([nombre, n])
    rs.column_dimensions["A"].width = 38; rs.column_dimensions["B"].width = 16
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


def pdf_bytes(rows, desde=None, hasta=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28,
                            title="Informe de reportes - PoliReport")
    st = getSampleStyleSheet()
    celda = ParagraphStyle("c", parent=st["BodyText"], fontSize=8, leading=10)
    verde = colors.HexColor("#173B36")
    res = resumen(rows)
    el = [Paragraph("Informe de reportes", st["Title"]),
          Paragraph(f"{escape(periodo(desde, hasta))} &nbsp;|&nbsp; Generado el {a_local(datetime.now(timezone.utc).strftime(FMT))}", st["Normal"]),
          Spacer(1, 10),
          Paragraph(f"<b>Total:</b> {res['total']} &nbsp;&nbsp; <b>Tiempo promedio de solución:</b> {escape(res['promedio'])}", st["Normal"]),
          Spacer(1, 8)]
    bloques = []
    for titulo, clave in (("Por estado", "por_estado"), ("Por urgencia", "por_prioridad"),
                          ("Por tipo", "por_tipo"), ("Áreas con más reportes", "por_area")):
        datos = [[titulo, "Reportes"]] + [[escape(str(a))[:34], n] for a, n in res[clave][:6]]
        t = Table(datos, colWidths=[140, 50])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), verde), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                               ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), .3, colors.lightgrey)]))
        bloques.append(t)
    el.append(Table([bloques], hAlign="LEFT"))
    el.append(Spacer(1, 12))
    cab = ["Fecha", "Área", "Tipo", "Urgencia", "Estado", "Descripción", "Encargado"]
    datos = [cab] + [[a_local(r["fecha"])[:10], Paragraph(escape(r["lugar"]), celda), Paragraph(escape(r["categoria"]), celda),
                      r["prioridad"], r["estado"],
                      Paragraph(escape(((r.get("problema") + ": ") if r.get("problema") else "") + r["descripcion"])[:220], celda),
                      Paragraph(escape(r.get("encargado") or "Sin asignar"), celda)] for r in rows]
    t = Table(datos, repeatRows=1, colWidths=[52, 120, 80, 50, 60, 330, 90])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), verde), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("FONTSIZE", (0, 0), (-1, -1), 8), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("GRID", (0, 0), (-1, -1), .3, colors.lightgrey),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F2")])]))
    el.append(t)
    doc.build(el)
    return out.getvalue()
