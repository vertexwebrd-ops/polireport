import io, os, tempfile
os.environ["DATABASE"] = os.path.join(tempfile.mkdtemp(), "f.db")
from app import app
import openpyxl

def cliente(correo, pw):
    c = app.test_client(); r = c.post("/api/login", json={"correo": correo, "password": pw}); assert r.status_code == 200, correo; return c

def nuevo_usuario(n):
    c = app.test_client()
    assert c.post("/api/registro", json={"nombre": f"Est {n}", "correo": f"e{n}@x.com", "password": "123456"}).status_code == 201
    return c

adm, est, mant = cliente("admin@poli.edu", "admin123"), cliente("estudiante@poli.edu", "demo123"), cliente("juan@poli.edu", "demo123")
rid = est.post("/api/reportes", json={"descripcion": "No enciende la PC", "id_lugar": 8, "id_categoria": 6, "prioridad": "Baja",
                                       "problema": "Computadoras que no funcionan"}).json["id_reporte"]
# aviso al admin por el reporte nuevo
av = adm.get("/api/avisos").json
assert av["no_leidos"] >= 1 and "Nuevo reporte" in av["items"][0]["texto"]
assert adm.get("/api/avisos/conteo").json["no_leidos"] == av["no_leidos"]
adm.post("/api/avisos/leer", json={}); assert adm.get("/api/avisos/conteo").json["no_leidos"] == 0

# similares + apoyos + escalado de urgencia (Baja -> Media con 3, -> Alta con 5)
otros = [nuevo_usuario(i) for i in range(5)]
sim = otros[0].get("/api/similares?id_lugar=8&id_categoria=6").json
assert len(sim) == 1 and sim[0]["apoyos"] == 0 and not sim[0]["es_mio"]
assert est.post(f"/api/reportes/{rid}/apoyo").status_code == 400          # no puede apoyar el suyo
for i, o in enumerate(otros[:3]):
    r = o.post(f"/api/reportes/{rid}/apoyo"); assert r.status_code == 200 and r.json["apoyos"] == i + 1
assert r.json["prioridad"] == "Media"
assert otros[0].post(f"/api/reportes/{rid}/apoyo").json["apoyos"] == 2     # quitar apoyo
otros[0].post(f"/api/reportes/{rid}/apoyo")
for o in otros[3:]: o.post(f"/api/reportes/{rid}/apoyo")
d = adm.get(f"/api/reportes/{rid}").json
assert d["apoyos"] == 5 and d["prioridad"] == "Alta"
assert any("personas reportaron" in e["texto"] for e in d["eventos"])
assert adm.post(f"/api/reportes/{rid}/apoyo").status_code == 403          # admin no apoya

# asignación -> aviso al encargado y al autor, y comentarios
adm.patch(f"/api/reportes/{rid}", json={"id_encargado": 1, "id_estado": 1})
assert "Te asignaron" in mant.get("/api/avisos").json["items"][0]["texto"]
assert any("asignado a Juan" in a["texto"] for a in est.get("/api/avisos").json["items"])
assert mant.post(f"/api/reportes/{rid}/comentarios", json={"texto": "Pieza pedida"}).status_code == 201
assert est.post(f"/api/reportes/{rid}/comentarios", json={"texto": "Gracias"}).status_code == 201
assert otros[1].post(f"/api/reportes/{rid}/comentarios", json={"texto": "hola"}).status_code == 404
tl = est.get(f"/api/reportes/{rid}").json["eventos"]
assert [e["tipo"] for e in tl if e["tipo"] in ("creado", "asignado", "comentario")] == ["creado", "asignado", "comentario", "comentario"]

# foto del arreglo y cierre
FOTO = "data:image/jpeg;base64,AAAA"
assert mant.patch(f"/api/reportes/{rid}", json={"id_estado": 3, "foto_despues": "x"}).status_code == 400
assert mant.patch(f"/api/reportes/{rid}", json={"id_estado": 3, "foto_despues": FOTO}).status_code == 200
assert est.get(f"/api/reportes/{rid}").json["foto_despues"] == FOTO
assert est.get("/api/reportes").json[0]["tiene_foto_despues"] and "foto_despues" not in est.get("/api/reportes").json[0]
assert any("pasó a Resuelto" in a["texto"] for a in est.get("/api/avisos").json["items"])
assert otros[0].post(f"/api/reportes/{rid}/apoyo").status_code == 404      # ya resuelto

# estadísticas
s = adm.get("/api/estadisticas").json
assert s["promedio_horas"] is not None and "atrasados" in s

# informes
x = adm.get("/api/informe.xlsx"); assert x.status_code == 200
wb = openpyxl.load_workbook(io.BytesIO(x.data)); assert wb.sheetnames == ["Reportes", "Resumen"] and wb["Reportes"].max_row == 2
assert adm.get("/api/informe.pdf").data[:4] == b"%PDF"
assert openpyxl.load_workbook(io.BytesIO(adm.get("/api/informe.xlsx?desde=2000-01-01&hasta=2000-01-02").data))["Reportes"].max_row == 1
assert est.get("/api/informe.pdf").status_code == 403

# gestión de encargados
r = adm.post("/api/admin/encargados", json={"nombre": "Ana Ruiz", "especialidad": "Plomería", "correo": "ana@poli.edu", "password": "clave123"})
assert r.status_code == 201; eid = r.json["id_encargado"]
assert cliente("ana@poli.edu", "clave123").get("/api/yo").json["rol"] == "mantenimiento"
assert adm.post("/api/admin/encargados", json={"nombre": "Otra", "correo": "ana@poli.edu", "password": "clave123"}).status_code == 409
assert adm.patch(f"/api/admin/encargados/{eid}", json={"nombre": "Ana R.", "especialidad": "Plomería"}).status_code == 200
assert adm.delete("/api/admin/encargados/1").status_code == 409            # tiene reportes
assert adm.delete(f"/api/admin/encargados/{eid}").status_code == 200
assert app.test_client().post("/api/login", json={"correo": "ana@poli.edu", "password": "clave123"}).status_code == 401
# gestión de áreas
lid = adm.post("/api/admin/lugares", json={"nombre": "Biblioteca", "grupo": "Áreas generales"}).json["id_lugar"]
assert adm.post("/api/admin/lugares", json={"nombre": "Biblioteca"}).status_code == 409
assert adm.patch(f"/api/admin/lugares/{lid}", json={"nombre": "Biblioteca central", "grupo": "Áreas generales"}).status_code == 200
assert adm.delete("/api/admin/lugares/8").status_code == 409                # en uso
assert adm.delete(f"/api/admin/lugares/{lid}").status_code == 200
assert est.post("/api/admin/lugares", json={"nombre": "X1"}).status_code == 403
print("FUNCIONES OK")
